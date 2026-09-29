"""S8 回答评测 runner（PLAN §5 S8）。

三个子命令：

- ``run``：对评测集的题逐题调 ``POST /v1/chat/stream``（真实 Agent + 真实 LLM，会花钱），保存每题的完整记录到
  ``<out>/<config>/answers.jsonl``（SSE 原文在 ``raw/``，不入库）。
  **只有 API 本身报错（超时 / 5xx / 连接错误）的题才重试，最多 2 次，每次尝试都记录**；其他失败（agent_error、
  没有回答、客户端超时……）保留为失败，绝不重跑到成功为止。已有记录的题在续跑时跳过。
- ``score``：读 answers.jsonl，规则判分 + LLM 裁判（``JUDGE_MODEL``），
  写 ``scores.jsonl`` 和总的 ``summary.json`` / ``report.md``。
- ``blind`` / ``blind-score``：生成人工盲标表（只有问题、参考要点、Agent 回答）/ 用户标完后算一致率与 kappa。

检索配置（``vector`` / ``hybrid_rerank``）由被测 ai-service 进程的 ``RETRIEVAL_MODE`` 决定（文档检索 MCP 用默认配置，
S5 起 Agent 不能改检索参数）；本脚本只负责在不同端口的服务上跑同一批题，并把端口对应的配置名记进 summary。
"""

from __future__ import annotations

import argparse
import asyncio
import json
import platform
import random
import sys
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from pathlib import Path

import httpx

from fund_ai.config import REPO_ROOT, get_settings
from fund_ai.eval.answer_score import (
    PRICES,
    ChunkIndex,
    aggregate,
    is_api_error,
    score_item,
)
from fund_ai.eval.retrieval import env_block, sha256
from fund_ai.eval.stats import paired_bootstrap

DATASETS = {"qa": "fund_qa_v1.jsonl", "agent": "agent_tasks_v1.jsonl"}
MAX_API_RETRIES = 2
CLIENT_TIMEOUT_S = 420.0
NAV_URL = "https://api.fund.eastmoney.com/f10/lsjz"


def load_items(split: str) -> list[dict]:
    """两个数据集里指定 split 的全部题（评测集校验与冻结见 S3）。核对数据集 sha256 与 MANIFEST 一致。"""
    manifest = json.loads((REPO_ROOT / "eval/datasets/MANIFEST.json").read_text(encoding="utf-8"))
    out = []
    for ds, fn in DATASETS.items():
        p = REPO_ROOT / "eval/datasets" / fn
        if sha256(p) != manifest["files"][fn]["sha256"]:
            raise SystemExit(f"{fn} 的 sha256 与 eval/datasets/MANIFEST.json 不一致，拒绝评测")
        for line in p.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            if split == "all" or r["split"] == split:
                r["_dataset"] = ds
                out.append(r)
    return out


def sample_per_topic(items: list[dict], n: int) -> list[dict]:
    """费用小样：每个 (数据集, topic) 取 id 最小的 n 题，确定性。"""
    by: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for r in sorted(items, key=lambda x: x["id"]):
        by[(r["_dataset"], r["topic"])].append(r)
    return [r for k in sorted(by) for r in by[k][:n]]


# ------------------------------------------------------------------ 参考净值（独立于被测工具的实现）
def fetch_reference_nav(share_code: str) -> dict | None:
    """volatile 题的标准答案：评测时直接请求东方财富净值接口（自己的代码路径，不经过被测的 MCP 工具），最多 3 次。"""
    for i in range(3):
        try:
            r = httpx.get(
                NAV_URL,
                params={"fundCode": share_code, "pageIndex": 1, "pageSize": 1},
                headers={
                    "Referer": "https://fundf10.eastmoney.com/",
                    "User-Agent": "fund-research-agent-eval/0.1",
                },
                timeout=8,
            )
            row = r.json()["Data"]["LSJZList"][0]
            return {
                "share_code": share_code,
                "unit_nav": float(row["DWJZ"]),
                "nav_date": row["FSRQ"],
                "fetched_at": datetime.now(UTC).isoformat(timespec="seconds"),
            }
        except Exception:  # noqa: BLE001
            time.sleep(1 + i)
    return None


# ------------------------------------------------------------------ 一次提问
async def ask_once(client: httpx.AsyncClient, url: str, question: str, raw_path: Path) -> dict:
    """调一次 /v1/chat/stream，解析 SSE。返回本次尝试的记录（不含重试逻辑）。"""
    events: list[tuple[str, dict]] = []
    t0 = time.perf_counter()
    err = ""
    try:
        async with asyncio.timeout(CLIENT_TIMEOUT_S):
            async with client.stream(
                "POST", f"{url}/v1/chat/stream", json={"question": question}
            ) as resp:
                resp.raise_for_status()
                cur = "message"
                raw_lines: list[str] = []
                async for line in resp.aiter_lines():
                    raw_lines.append(line)
                    if line.startswith("event:"):
                        cur = line[6:].strip()
                    elif line.startswith("data:"):
                        events.append((cur, json.loads(line[5:].strip())))
                raw_path.parent.mkdir(parents=True, exist_ok=True)
                raw_path.write_text("\n".join(raw_lines) + "\n", encoding="utf-8", newline="\n")
    except TimeoutError:
        err = f"client_timeout: 超过 {CLIENT_TIMEOUT_S:.0f}s"
    except Exception as e:  # noqa: BLE001
        err = f"client_error: {type(e).__name__}: {str(e)[:200]}"
    wall_ms = round((time.perf_counter() - t0) * 1000, 1)

    answer = "".join(d.get("text", "") for e, d in events if e == "token")
    done = next((d for e, d in reversed(events) if e == "done"), {})
    agent_err = next((d for e, d in events if e == "error"), None)
    calls: dict[str, dict] = {}
    for e, d in events:
        if e == "tool_start":
            calls[d["call_id"]] = {"name": d["name"], "args": d.get("args"), "step": d.get("step")}
        elif e == "tool_end" and d.get("call_id") in calls:
            calls[d["call_id"]].update(
                {
                    k: d.get(k)
                    for k in (
                        "status",
                        "duration_ms",
                        "summary",
                        "citation_ids",
                        "error",
                        "error_kind",
                    )
                }
            )
    types: list[str] = []
    for e, _ in events:
        if e != "token" or not types or types[-1] != "token":
            types.append(e)
    status = "ok" if (done.get("status") == "ok" and answer.strip() and not err) else "error"
    if not err and agent_err:
        err = f"{agent_err.get('code')}: {agent_err.get('message')}"
    elif not err and status == "error":
        err = f"done.status={done.get('status')!r}, answer_chars={len(answer)}"
    return {
        "status": status,
        "error": err or None,
        "answer": answer,
        "done": done,
        "tools": list(calls.values()),
        "citations": next((d.get("items", []) for e, d in events if e == "citations"), []),
        "event_types": types,
        "disclaimer_text": next((d.get("text") for e, d in events if e == "disclaimer"), None),
        "wall_ms": wall_ms,
    }


async def ask_with_retry(client: httpx.AsyncClient, url: str, item: dict, raw_dir: Path) -> dict:
    attempts: list[dict] = []
    for n in range(1, MAX_API_RETRIES + 2):
        res = await ask_once(client, url, item["question"], raw_dir / f"{item['id']}.a{n}.sse")
        attempts.append(
            {"n": n, "status": res["status"], "error": res["error"], "wall_ms": res["wall_ms"]}
        )
        if res["status"] == "ok" or not is_api_error(res["error"] or "") or n > MAX_API_RETRIES:
            break
    rec = {
        "id": item["id"],
        "dataset": item["_dataset"],
        "split": item["split"],
        "topic": item["topic"],
        "attempts": attempts,
        "retried_for_api_error": len(attempts) > 1,
        "final": {k: res[k] for k in ("status", "error", "answer", "done")},
        "tools": res["tools"],
        "citations": res["citations"],
        "event_types": res["event_types"],
        "disclaimer_text": res["disclaimer_text"],
    }
    if item.get("volatile") and res["status"] == "ok":
        rec["reference_nav"] = await asyncio.to_thread(
            fetch_reference_nav, (item.get("gold_params") or {})["share_code"]
        )
    return rec


async def run_all(args: argparse.Namespace, items: list[dict], out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    ans_path = out / "answers.jsonl"
    done_ids = set()
    if ans_path.exists():
        done_ids = {
            json.loads(line)["id"] for line in ans_path.read_text(encoding="utf-8").splitlines()
        }
    todo = [it for it in items if it["id"] not in done_ids]
    print(
        f"[{args.config}] {len(items)} 题，已有记录 {len(done_ids)}，本次运行 {len(todo)}",
        flush=True,
    )
    sem = asyncio.Semaphore(args.concurrency)
    lock = asyncio.Lock()
    finished = 0
    async with httpx.AsyncClient(
        timeout=httpx.Timeout(CLIENT_TIMEOUT_S, connect=10), trust_env=False
    ) as client:

        async def one(item: dict) -> None:
            nonlocal finished
            async with sem:
                rec = await ask_with_retry(client, args.url, item, out / "raw")
            async with lock:
                with ans_path.open("a", encoding="utf-8", newline="\n") as f:
                    f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                finished += 1
                tag = "ok" if rec["final"]["status"] == "ok" else "FAILED"
                print(
                    f"  ({finished}/{len(todo)}) {item['id']} {tag} attempts={len(rec['attempts'])}",
                    flush=True,
                )

        await asyncio.gather(*(one(it) for it in todo))


def cmd_run(a: argparse.Namespace) -> int:
    items = load_items(a.split)
    if a.sample_per_topic:
        items = sample_per_topic(items, a.sample_per_topic)
    if a.only:
        want = set(a.only.split(","))
        items = [r for r in items if r["id"] in want]
    out_root = (
        Path(a.out)
        if a.out
        else REPO_ROOT / "reports" / "answer_eval" / datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    )
    out = out_root / a.config
    asyncio.run(run_all(a, items, out))
    meta = {
        "config": a.config,
        "service_url": a.url,
        "split": a.split,
        "sample_per_topic": a.sample_per_topic,
        "n_items": len(items),
        "concurrency": a.concurrency,
        "command": "cd ai-service && python -m fund_ai.eval.answer " + " ".join(sys.argv[1:]),
        "finished_utc": datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ"),
    }
    (out / "run_meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print(f"结果目录：{out_root}")
    return 0


# ------------------------------------------------------------------ 判分与汇总
def _load_jsonl(p: Path) -> list[dict]:
    return [json.loads(line) for line in p.read_text(encoding="utf-8").splitlines() if line.strip()]


def cmd_score(a: argparse.Namespace) -> int:
    from fund_ai.eval.judge import Judge

    root = Path(a.dir)
    settings = get_settings()
    judge = Judge(settings)
    chunks = ChunkIndex(REPO_ROOT / "data" / "raw" / "_chunks.jsonl")
    all_items = {r["id"]: r for r in load_items("all")}
    configs = [d.name for d in sorted(root.iterdir()) if (d / "answers.jsonl").exists()]
    per_cfg: dict[str, dict] = {}
    for cfg in configs:
        recs = {r["id"]: r for r in _load_jsonl(root / cfg / "answers.jsonl")}
        sp = root / cfg / "scores.jsonl"
        have = {r["id"]: r for r in _load_jsonl(sp)} if sp.exists() and not a.rescore else {}
        todo = [i for i in recs if i not in have]
        print(f"[{cfg}] 判分 {len(todo)} 题（已有 {len(have)}）", flush=True)
        with ThreadPoolExecutor(max_workers=a.judge_concurrency) as ex:
            futs = {i: ex.submit(score_item, all_items[i], recs[i], judge, chunks) for i in todo}
            for i, f in futs.items():
                have[i] = f.result()
        order = sorted(have)
        with sp.open("w", encoding="utf-8", newline="\n") as f:
            for i in order:
                f.write(json.dumps(have[i], ensure_ascii=False) + "\n")
        per_cfg[cfg] = {"recs": recs, "scores": [have[i] for i in order]}

    results = {cfg: aggregate(v["scores"], v["recs"]) for cfg, v in per_cfg.items()}
    comparison = compare_configs(per_cfg) if len(per_cfg) == 2 else {}
    metas = {
        cfg: json.loads((root / cfg / "run_meta.json").read_text(encoding="utf-8"))
        for cfg in configs
        if (root / cfg / "run_meta.json").exists()
    }
    retries = {
        cfg: [
            {"id": r["id"], "attempts": r["attempts"]}
            for r in v["recs"].values()
            if r["retried_for_api_error"]
        ]
        for cfg, v in per_cfg.items()
    }
    failed = {
        cfg: [
            {"id": r["id"], "error": r["final"]["error"], "attempts": len(r["attempts"])}
            for r in v["recs"].values()
            if r["final"]["status"] != "ok"
        ]
        for cfg, v in per_cfg.items()
    }
    resp_models = sorted(
        {
            m
            for v in per_cfg.values()
            for r in v["recs"].values()
            for m in r["final"]["done"].get("response_models", [])
        }
    )
    req_models = sorted(
        {
            r["final"]["done"].get("request_model")
            for v in per_cfg.values()
            for r in v["recs"].values()
            if r["final"]["done"].get("request_model")
        }
    )
    judge_resp = sorted(
        {
            ((s.get("judge") or {}).get("call") or {}).get("response_model", "")
            for v in per_cfg.values()
            for s in v["scores"]
        }
        - {""}
    )
    env = env_block(a.ingest_run)
    env["models"].update(
        llm_requested=req_models,
        llm_response=resp_models,
        judge_requested=settings.judge_model,
        judge_response=judge_resp,
    )
    env["dataset_sha256"] = {
        fn: sha256(REPO_ROOT / "eval/datasets" / fn) for fn in DATASETS.values()
    }
    summary = {
        "label": a.label or "S8 回答评测",
        "env": env,
        "params": {
            "agent_max_steps": settings.agent_max_steps,
            "agent_sql_retries": settings.agent_sql_retries,
            "llm_thinking": settings.llm_thinking,
            "api_retries_max": MAX_API_RETRIES,
            "numeric_rule": "见 fund_ai/eval/scoring.py 文件头",
            "prices": PRICES,
        },
        "runs": metas,
        "results": results,
        "comparison_vector_to_hybrid_rerank": comparison,
        "api_error_retries": retries,
        "failed_runs": failed,
        "machine": {"platform": platform.platform(), "python": sys.version.split()[0]},
        "scored_at_utc": datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ"),
        "score_command": "cd ai-service && python -m fund_ai.eval.answer " + " ".join(sys.argv[1:]),
    }
    (root / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    from fund_ai.eval.answer_report import render

    (root / "report.md").write_text(render(summary), encoding="utf-8", newline="\n")
    print(f"summary → {root / 'summary.json'}")
    return 0


def compare_configs(per_cfg: dict[str, dict]) -> dict:
    """vector → hybrid_rerank 的配对 bootstrap（逐题统一得分 value，只取两边都有得分的题）。"""
    if not {"vector", "hybrid_rerank"} <= set(per_cfg):
        return {}
    a = {s["id"]: s for s in per_cfg["vector"]["scores"]}
    b = {s["id"]: s for s in per_cfg["hybrid_rerank"]["scores"]}
    ids = sorted(i for i in a if i in b and a[i]["value"] is not None and b[i]["value"] is not None)

    def boot(sel: list[str]) -> dict:
        return paired_bootstrap([a[i]["value"] for i in sel], [b[i]["value"] for i in sel])

    def used_docs(i: str, cfg: str) -> bool:
        return any(
            t["name"] == "search_fund_documents" for t in per_cfg[cfg]["recs"][i].get("tools") or []
        )

    docs_ids = [i for i in ids if used_docs(i, "vector") or used_docs(i, "hybrid_rerank")]
    by_topic = {}
    for t in sorted({a[i]["topic"] for i in ids}):
        sel = [i for i in ids if a[i]["topic"] == t]
        by_topic[t] = boot(sel)
    return {
        "unit": "逐题统一得分 value（规则类 0/1，text 类裁判分/2，refusal 类 0/1）；差 = hybrid_rerank − vector",
        "all_questions": boot(ids),
        "questions_where_docs_tool_used_in_either_run": boot(docs_ids),
        "by_topic": by_topic,
    }


# ------------------------------------------------------------------ 盲标
def cmd_blind(a: argparse.Namespace) -> int:
    root = Path(a.dir)
    all_items = {r["id"]: r for r in load_items("all")}
    pool = []
    for cfg in sorted(d.name for d in root.iterdir() if (d / "answers.jsonl").exists()):
        for r in _load_jsonl(root / cfg / "answers.jsonl"):
            it = all_items[r["id"]]
            if it["answer_type"] == "text" and r["final"]["status"] == "ok":
                pool.append((cfg, it, r))
    by_topic: dict[str, list] = defaultdict(list)
    for x in sorted(pool, key=lambda t: (t[1]["topic"], t[1]["id"], t[0])):
        by_topic[x[1]["topic"]].append(x)
    rng = random.Random(a.seed)
    total = len(pool)
    picked = []
    for xs in (by_topic[t] for t in sorted(by_topic)):
        k = max(1, round(a.n * len(xs) / total))
        picked += rng.sample(xs, min(k, len(xs)))
    rng.shuffle(picked)
    rows, key = [], []
    for i, (cfg, it, r) in enumerate(picked, 1):
        rid = f"B{i:03d}"
        rows.append(
            {
                "row": rid,
                "question": it["question"],
                "points": it["answer_points"],
                "answer": r["final"]["answer"],
            }
        )
        key.append({"row": rid, "config": cfg, "id": it["id"], "topic": it["topic"]})
    (root / "blind_key.json").write_text(
        json.dumps(
            {"seed": a.seed, "n": len(key), "pool": total, "rows": key},
            ensure_ascii=False,
            indent=1,
        )
        + "\n",
        encoding="utf-8",
        newline="\n",
    )
    write_blind_table(rows, Path(a.out))
    print(
        f"盲标表 {len(rows)} 行（候选池 {total}，种子 {a.seed}）→ {a.out}；对照表 {root / 'blind_key.json'}"
    )
    return 0


HEADER_TEXT = (
    "评分标准：0 = 错误或缺失关键要点；1 = 部分正确；2 = 完整正确。请只根据「参考要点」判断回答的事实是否正确，"
    "不要看文采和篇幅；回答比要点多出的、不矛盾的信息不扣分。请在「你的评分」列填 0、1 或 2。"
)


def write_blind_table(rows: list[dict], out: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:  # 没有 openpyxl 时退回 CSV
        import csv

        with out.with_suffix(".csv").open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow([HEADER_TEXT])
            w.writerow(["行号", "问题", "参考要点", "Agent 回答", "你的评分（0/1/2）"])
            for r in rows:
                w.writerow(
                    [
                        r["row"],
                        r["question"],
                        "\n".join(f"{i}. {p}" for i, p in enumerate(r["points"], 1)),
                        r["answer"],
                        "",
                    ]
                )
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "盲标"
    ws.append([HEADER_TEXT])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True)
    ws.row_dimensions[1].height = 60
    ws.append(["行号", "问题", "参考要点", "Agent 回答", "你的评分（0/1/2）"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                r["row"],
                r["question"],
                "\n".join(f"{i}. {p}" for i, p in enumerate(r["points"], 1)),
                r["answer"],
                None,
            ]
        )
    for col, w in zip("ABCDE", (8, 40, 50, 80, 14), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    wb.save(out.with_suffix(".xlsx"))


def cohen_kappa(x: list[int], y: list[int], weights: str | None = None) -> float:
    cats = sorted(set(x) | set(y) | {0, 1, 2})
    n = len(x)
    idx = {c: i for i, c in enumerate(cats)}
    k = len(cats)
    obs = [[0.0] * k for _ in range(k)]
    for p, q in zip(x, y, strict=True):
        obs[idx[p]][idx[q]] += 1
    row = [sum(r) for r in obs]
    col = [sum(obs[i][j] for i in range(k)) for j in range(k)]

    def w(i: int, j: int) -> float:
        if weights == "linear":
            return abs(i - j) / (k - 1)
        return 0.0 if i == j else 1.0

    num = sum(w(i, j) * obs[i][j] for i in range(k) for j in range(k)) / n
    den = sum(w(i, j) * row[i] * col[j] for i in range(k) for j in range(k)) / (n * n)
    return 1.0 if den == 0 else 1 - num / den


def cmd_blind_score(a: argparse.Namespace) -> int:
    root = Path(a.dir)
    key = {
        r["row"]: r
        for r in json.loads((root / "blind_key.json").read_text(encoding="utf-8"))["rows"]
    }
    from openpyxl import load_workbook

    ws = load_workbook(a.table, data_only=True).active
    human: dict[str, int] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] and row[4] is not None and str(row[4]).strip() != "":
            human[str(row[0])] = int(float(str(row[4]).strip()))
    scores = {}
    for cfg in {v["config"] for v in key.values()}:
        for s in _load_jsonl(root / cfg / "scores.jsonl"):
            scores[(cfg, s["id"])] = s
    pairs = []
    for rid, h in sorted(human.items()):
        k = key[rid]
        s = scores[(k["config"], k["id"])]
        if s["correct"] is None:
            continue
        pairs.append((rid, h, s["detail"]["score"]))
    hs, js = [p[1] for p in pairs], [p[2] for p in pairs]
    n = len(pairs)
    conf = [
        [sum(1 for x, y in zip(hs, js, strict=True) if x == i and y == j) for j in range(3)]
        for i in range(3)
    ]
    out = {
        "n_key_rows": len(key),
        "n_human_labelled": len(human),
        "n_compared": n,
        "n_judge_failed_excluded": len(human) - n,
        "exact_agreement": round(sum(x == y for x, y in zip(hs, js, strict=True)) / n, 4),
        "cohen_kappa_unweighted": round(cohen_kappa(hs, js), 4),
        "cohen_kappa_linear_weighted": round(cohen_kappa(hs, js, "linear"), 4),
        "binary_full_score_agreement": round(
            sum((x == 2) == (y == 2) for x, y in zip(hs, js, strict=True)) / n, 4
        ),
        "binary_kappa_full_vs_not": round(
            cohen_kappa([int(x == 2) for x in hs], [int(y == 2) for y in js]), 4
        ),
        "confusion_rows_human_cols_judge": conf,
        "human_mean_score": round(sum(hs) / n, 4),
        "judge_mean_score": round(sum(js) / n, 4),
        "pairs": [
            {
                "row": r,
                **{"config": key[r]["config"], "id": key[r]["id"], "topic": key[r]["topic"]},
                "human": h,
                "judge": j,
            }
            for r, h, j in pairs
        ],
    }
    (root / "blind_agreement.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    print({k: v for k, v in out.items() if k != "pairs"})
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--config", required=True, choices=["vector", "hybrid_rerank"])
    r.add_argument("--url", required=True, help="RETRIEVAL_MODE 与 --config 一致的 ai-service 地址")
    r.add_argument("--split", choices=["dev", "test", "all"], default="test")
    r.add_argument(
        "--sample-per-topic",
        type=int,
        default=0,
        help="每个 (数据集, topic) 取前 N 题（费用小样，用 dev）",
    )
    r.add_argument("--only", help="逗号分隔的题 id")
    r.add_argument(
        "--out", help="结果根目录，缺省 reports/answer_eval/<UTC>；续跑与第二个配置要传同一个"
    )
    r.add_argument("--concurrency", type=int, default=3)
    r.set_defaults(fn=cmd_run)
    s = sub.add_parser("score")
    s.add_argument("--dir", required=True)
    s.add_argument("--label", default="")
    s.add_argument("--ingest-run")
    s.add_argument("--judge-concurrency", type=int, default=4)
    s.add_argument("--rescore", action="store_true")
    s.set_defaults(fn=cmd_score)
    b = sub.add_parser("blind")
    b.add_argument("--dir", required=True)
    b.add_argument("--out", required=True, help="输出路径（不含扩展名）")
    b.add_argument("--n", type=int, default=36)
    b.add_argument("--seed", type=int, default=20260930)
    b.set_defaults(fn=cmd_blind)
    bs = sub.add_parser("blind-score")
    bs.add_argument("--dir", required=True)
    bs.add_argument("--table", required=True)
    bs.set_defaults(fn=cmd_blind_score)
    a = ap.parse_args(argv)
    return a.fn(a)


if __name__ == "__main__":
    raise SystemExit(main())
